"""Loading sessions in the database: saving what workers report, and the dashboard's queries."""

import io
from collections import defaultdict
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import UTC, date, datetime, timedelta
from zoneinfo import ZoneInfo

from sqlalchemy import Select, func, select
from sqlalchemy.ext.asyncio import AsyncSession

from yardlens_api.core.enums import SessionStatus
from yardlens_api.core.i18n import EXPORT_LABELS, Language
from yardlens_api.db.models import Camera, LoadingSession
from yardlens_api.schemas.loading import (
    DailyLoading,
    LoadingSessionPage,
    LoadingSessionRead,
    LoadingSummary,
    LoadingTotals,
)
from yardlens_api.services.events import snapshot_url
from yardlens_api.services.loading import SessionState
from yardlens_api.services.plates import format_plate, normalize_plate

# A stored row or a worker's in-memory session; both carry the same fields.
type SessionRecord = LoadingSession | SessionState

# A session saved this recently when the API stopped is resumed rather than closed: the vehicle is
# probably still in the bay. Open sessions are saved at least once a minute while it is.
RESUME_WINDOW = timedelta(minutes=5)


@dataclass(frozen=True, slots=True)
class InterruptedSession:
    """An open session found at startup, handed to its camera's worker to carry on."""

    state: SessionState
    last_seen: datetime


def to_session_read(
    record: SessionRecord, camera_name: str, *, corrected: bool = False
) -> LoadingSessionRead:
    end = record.departed_at or datetime.now(UTC)
    return LoadingSessionRead(
        id=record.id,
        camera_id=record.camera_id,
        camera_name=camera_name,
        status=SessionStatus(record.status),
        arrived_at=record.arrived_at,
        departed_at=record.departed_at,
        loading_started_at=record.loading_started_at,
        loading_finished_at=record.loading_finished_at,
        duration_seconds=max(0, round((end - record.arrived_at).total_seconds())),
        loaded=record.loaded,
        unloaded=record.unloaded,
        vehicle_class=record.vehicle_class,
        plate=record.plate,
        plate_confidence=record.plate_confidence,
        plate_corrected=corrected,
        vehicle_snapshot_url=snapshot_url(record.vehicle_snapshot_path),
        plate_snapshot_url=snapshot_url(record.plate_snapshot_path),
    )


def row_to_read(row: LoadingSession, camera_name: str) -> LoadingSessionRead:
    return to_session_read(row, camera_name, corrected=row.plate_corrected)


async def save_session(session: AsyncSession, state: SessionState) -> LoadingSession:
    """Insert or update the row for a worker's session. A plate fixed by hand is kept."""
    row = await session.get(LoadingSession, state.id)
    if row is None:
        row = LoadingSession(id=state.id, camera_id=state.camera_id, arrived_at=state.arrived_at)
        session.add(row)
    row.status = state.status.value
    row.arrived_at = state.arrived_at
    row.departed_at = state.departed_at
    row.loading_started_at = state.loading_started_at
    row.loading_finished_at = state.loading_finished_at
    row.loaded = state.loaded
    row.unloaded = state.unloaded
    row.vehicle_class = state.vehicle_class
    if not row.plate_corrected:
        row.plate = state.plate
        row.plate_confidence = state.plate_confidence
    row.vehicle_snapshot_path = state.vehicle_snapshot_path or row.vehicle_snapshot_path
    row.plate_snapshot_path = state.plate_snapshot_path or row.plate_snapshot_path
    # Explicit: a heartbeat changes no other column, and it is what marks the vehicle as present.
    row.updated_at = datetime.now(UTC)
    await session.commit()
    return row


def row_to_state(row: LoadingSession) -> SessionState:
    return SessionState(
        id=row.id,
        camera_id=row.camera_id,
        arrived_at=row.arrived_at,
        status=SessionStatus(row.status),
        departed_at=row.departed_at,
        loading_started_at=row.loading_started_at,
        loading_finished_at=row.loading_finished_at,
        loaded=row.loaded,
        unloaded=row.unloaded,
        vehicle_class=row.vehicle_class,
        plate=row.plate,
        plate_confidence=row.plate_confidence,
        vehicle_snapshot_path=row.vehicle_snapshot_path,
        plate_snapshot_path=row.plate_snapshot_path,
    )


async def take_interrupted_sessions(
    session: AsyncSession, *, running_cameras: set[int], now: datetime | None = None
) -> tuple[dict[int, InterruptedSession], int]:
    """Sessions still open at startup were cut off by a restart. The newest one per running
    camera, if saved within ``RESUME_WINDOW``, is returned to be resumed; the rest are closed as
    of their last save. Returns the resumable sessions by camera and how many were closed."""
    now = now or datetime.now(UTC)
    rows = (
        await session.scalars(
            select(LoadingSession)
            .where(LoadingSession.status == SessionStatus.ACTIVE.value)
            .order_by(LoadingSession.updated_at.desc())
        )
    ).all()
    resumable: dict[int, InterruptedSession] = {}
    closed = 0
    for row in rows:
        if (
            row.camera_id in running_cameras
            and row.camera_id not in resumable
            and now - row.updated_at <= RESUME_WINDOW
        ):
            resumable[row.camera_id] = InterruptedSession(row_to_state(row), row.updated_at)
            continue
        row.status = SessionStatus.COMPLETED.value
        row.departed_at = max(row.updated_at, row.arrived_at)
        closed += 1
    await session.commit()
    return resumable, closed


@dataclass(frozen=True, slots=True)
class SessionFilters:
    camera_id: int | None = None
    plate: str | None = None
    since: datetime | None = None
    until: datetime | None = None


def _filtered[*Ts](statement: Select[*Ts], filters: SessionFilters) -> Select[*Ts]:
    if filters.camera_id is not None:
        statement = statement.where(LoadingSession.camera_id == filters.camera_id)
    if filters.plate:
        statement = statement.where(LoadingSession.plate.contains(normalize_plate(filters.plate)))
    if filters.since is not None:
        statement = statement.where(LoadingSession.arrived_at >= filters.since)
    if filters.until is not None:
        statement = statement.where(LoadingSession.arrived_at < filters.until)
    return statement


type _TotalsRow = tuple[int, int, datetime, datetime | None]


def _totals(rows: Sequence[_TotalsRow]) -> LoadingTotals:
    durations = [(left - came).total_seconds() for _, _, came, left in rows if left is not None]
    return LoadingTotals(
        vehicles=len(rows),
        loaded=sum(loaded for loaded, _, _, _ in rows),
        unloaded=sum(unloaded for _, unloaded, _, _ in rows),
        average_duration_seconds=round(sum(durations) / len(durations)) if durations else None,
    )


async def _totals_rows(session: AsyncSession, filters: SessionFilters) -> list[_TotalsRow]:
    statement = _filtered(
        select(
            LoadingSession.loaded,
            LoadingSession.unloaded,
            LoadingSession.arrived_at,
            LoadingSession.departed_at,
        ),
        filters,
    )
    return [
        (loaded, unloaded, came, left)
        for loaded, unloaded, came, left in await session.execute(statement)
    ]


async def list_sessions(
    session: AsyncSession, filters: SessionFilters, *, page: int, page_size: int
) -> LoadingSessionPage:
    statement = _filtered(select(LoadingSession, Camera.name).join(Camera), filters)
    total = await session.scalar(select(func.count()).select_from(statement.subquery())) or 0
    rows = await session.execute(
        statement.order_by(LoadingSession.arrived_at.desc())
        .offset((page - 1) * page_size)
        .limit(page_size)
    )
    return LoadingSessionPage(
        items=[row_to_read(row, name) for row, name in rows],
        total=total,
        page=page,
        page_size=page_size,
        totals=_totals(await _totals_rows(session, filters)),
    )


async def all_sessions(
    session: AsyncSession, filters: SessionFilters, *, limit: int
) -> list[LoadingSessionRead]:
    statement = _filtered(select(LoadingSession, Camera.name).join(Camera), filters)
    rows = await session.execute(statement.order_by(LoadingSession.arrived_at).limit(limit))
    return [row_to_read(row, name) for row, name in rows]


async def get_session_row(
    session: AsyncSession, session_id: str
) -> tuple[LoadingSession, str] | None:
    row = (
        await session.execute(
            select(LoadingSession, Camera.name).join(Camera).where(LoadingSession.id == session_id)
        )
    ).first()
    return (row[0], row[1]) if row is not None else None


def _local_midnight(day: date, tz: ZoneInfo) -> datetime:
    return datetime(day.year, day.month, day.day, tzinfo=tz).astimezone(UTC)


async def summary(
    session: AsyncSession, *, tz: ZoneInfo, days: int, now: datetime | None = None
) -> LoadingSummary:
    now = now or datetime.now(UTC)
    today = now.astimezone(tz).date()
    first_day = today - timedelta(days=days - 1)
    rows = await _totals_rows(session, SessionFilters(since=_local_midnight(first_day, tz)))

    by_day: dict[date, list[_TotalsRow]] = defaultdict(list)
    for row in rows:
        by_day[row[2].astimezone(tz).date()].append(row)
    daily = [
        DailyLoading(
            day=day,
            vehicles=len(by_day[day]),
            loaded=sum(loaded for loaded, _, _, _ in by_day[day]),
        )
        for day in (first_day + timedelta(days=offset) for offset in range(days))
    ]

    active_rows = await session.execute(
        select(LoadingSession, Camera.name)
        .join(Camera)
        .where(LoadingSession.status == SessionStatus.ACTIVE.value)
        .order_by(LoadingSession.arrived_at)
    )
    loading_cameras = await session.scalar(
        select(func.count()).select_from(Camera).where(Camera.zone.is_not(None))
    )
    return LoadingSummary(
        timezone=tz.key,
        loading_cameras=loading_cameras or 0,
        today=_totals(by_day[today]),
        active=[row_to_read(row, name) for row, name in active_rows],
        daily=daily,
    )


async def purge_session_images_before(session: AsyncSession, cutoff: datetime) -> list[str]:
    """Forget the photos of sessions that arrived before ``cutoff``; the sessions stay. Returns
    the files to delete."""
    rows = (
        await session.scalars(
            select(LoadingSession).where(
                LoadingSession.arrived_at < cutoff,
                (LoadingSession.vehicle_snapshot_path.is_not(None))
                | (LoadingSession.plate_snapshot_path.is_not(None)),
            )
        )
    ).all()
    paths: list[str] = []
    for row in rows:
        paths += [path for path in (row.vehicle_snapshot_path, row.plate_snapshot_path) if path]
        row.vehicle_snapshot_path = None
        row.plate_snapshot_path = None
    await session.commit()
    return paths


async def session_snapshot_paths(session: AsyncSession, camera_id: int) -> list[str]:
    rows = await session.execute(
        select(LoadingSession.vehicle_snapshot_path, LoadingSession.plate_snapshot_path).where(
            LoadingSession.camera_id == camera_id
        )
    )
    return [path for pair in rows for path in pair if path]


def build_workbook(
    sessions: Sequence[LoadingSessionRead], *, tz: ZoneInfo, language: Language
) -> bytes:
    """The sessions as an Excel sheet with a totals row, times in ``tz``."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    labels = EXPORT_LABELS[language]
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = labels["sheet"]
    columns = [
        ("number", 6),
        ("date", 12),
        ("plate", 16),
        ("camera", 22),
        ("arrived", 11),
        ("loading_started", 17),
        ("loading_finished", 17),
        ("departed", 11),
        ("minutes", 18),
        ("loaded", 12),
        ("unloaded", 12),
    ]
    sheet.append([labels[key] for key, _ in columns])
    for index, (_, width) in enumerate(columns, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.freeze_panes = "A2"

    def clock(moment: datetime | None) -> str:
        return moment.astimezone(tz).strftime("%H:%M") if moment else ""

    for number, item in enumerate(sessions, start=1):
        sheet.append(
            [
                number,
                item.arrived_at.astimezone(tz).date(),
                format_plate(item.plate) if item.plate else "-",
                item.camera_name,
                clock(item.arrived_at),
                clock(item.loading_started_at),
                clock(item.loading_finished_at),
                clock(item.departed_at) or labels["in_progress"],
                round(item.duration_seconds / 60),
                item.loaded,
                item.unloaded,
            ]
        )
        sheet.cell(row=sheet.max_row, column=2).number_format = "DD.MM.YYYY"

    sheet.append(
        [
            labels["total"],
            None,
            f"{len(sessions)} {labels['vehicles']}",
            None,
            None,
            None,
            None,
            None,
            None,
            sum(item.loaded for item in sessions),
            sum(item.unloaded for item in sessions),
        ]
    )
    for cell in sheet[sheet.max_row]:
        cell.font = Font(bold=True)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


async def resumable_session(
    session: AsyncSession, camera_id: int, *, now: datetime | None = None
) -> InterruptedSession | None:
    """The camera's open session, when it was saved recently enough to be carried on by the node
    the camera moves to."""
    now = now or datetime.now(UTC)
    row = await session.scalar(
        select(LoadingSession)
        .where(
            LoadingSession.camera_id == camera_id,
            LoadingSession.status == SessionStatus.ACTIVE.value,
        )
        .order_by(LoadingSession.updated_at.desc())
        .limit(1)
    )
    if row is None or now - row.updated_at > RESUME_WINDOW:
        return None
    return InterruptedSession(row_to_state(row), row.updated_at)
