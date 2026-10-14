import asyncio
import io
from dataclasses import replace
from datetime import UTC, datetime, timedelta
from typing import Any

import httpx
import pytest
from openpyxl import load_workbook

from yardlens_api.api.deps import Container
from yardlens_api.core.enums import SessionStatus
from yardlens_api.services.loading import SessionState

ZONE = {"points": [{"x": 0.0, "y": 0.5}, {"x": 1.0, "y": 0.5}, {"x": 0.5, "y": 1.0}]}


async def _wait_for_total(client: httpx.AsyncClient, total: int) -> dict[str, Any]:
    for _ in range(100):
        body: dict[str, Any] = (await client.get("/api/loading/sessions")).json()
        if body["total"] == total:
            return body
        await asyncio.sleep(0.05)
    raise AssertionError(f"expected {total} sessions")


@pytest.mark.usefixtures("video_file")
async def test_loading_sessions_are_saved_summarised_corrected_and_exported(
    client: httpx.AsyncClient, container: Container
) -> None:
    created = await client.post(
        "/api/cameras",
        json={
            "name": "Склад",
            "source_type": "file",
            "url": "videos/clip.mp4",
            "enabled": False,
            "zone": ZONE,
        },
    )
    camera_id = created.json()["id"]
    now = datetime.now(UTC)
    finished = SessionState(
        id="a" * 32,
        camera_id=camera_id,
        arrived_at=now - timedelta(minutes=40),
        status=SessionStatus.COMPLETED,
        departed_at=now - timedelta(minutes=10),
        loading_started_at=now - timedelta(minutes=35),
        loading_finished_at=now - timedelta(minutes=12),
        loaded=400,
        vehicle_class="truck",
        plate="01A123BC",
        plate_confidence=0.93,
    )
    loading = SessionState(
        id="b" * 32,
        camera_id=camera_id,
        arrived_at=now - timedelta(minutes=5),
        loaded=50,
        plate="01B777CD",
    )
    container.recorder.submit_session(finished)
    container.recorder.submit_session(loading)
    page = await _wait_for_total(client, 2)
    assert page["totals"]["loaded"] == 450
    assert page["totals"]["average_duration_seconds"] == 30 * 60
    assert [item["plate"] for item in page["items"]] == ["01B777CD", "01A123BC"]

    summary = (await client.get("/api/loading/summary", params={"tz": "UTC", "days": 7})).json()
    arrived_today = [s for s in (finished, loading) if s.arrived_at.date() == now.date()]
    assert summary["loading_cameras"] == 1
    assert summary["today"]["vehicles"] == len(arrived_today)
    assert len(summary["daily"]) == 7
    assert summary["daily"][-1]["loaded"] == sum(s.loaded for s in arrived_today)
    assert [item["plate"] for item in summary["active"]] == ["01B777CD"]

    found = (await client.get("/api/loading/sessions", params={"plate": "01 a 123"})).json()
    assert [item["id"] for item in found["items"]] == [finished.id]

    corrected = await client.patch(
        f"/api/loading/sessions/{finished.id}", json={"plate": "01 A 124 BC"}
    )
    assert corrected.json()["plate"] == "01A124BC"
    assert corrected.json()["plate_corrected"] is True
    # A later reading from the camera does not undo the correction.
    container.recorder.submit_session(replace(finished, loaded=401))
    session: dict[str, Any] = {}
    for _ in range(100):
        session = (await client.get(f"/api/loading/sessions/{finished.id}")).json()
        if session["loaded"] == 401:
            break
        await asyncio.sleep(0.05)
    assert session["loaded"] == 401
    assert session["plate"] == "01A124BC"

    export = await client.get("/api/loading/sessions/export", params={"tz": "UTC", "lang": "ru"})
    assert export.headers["content-type"].startswith("application/vnd.openxmlformats")
    sheet = load_workbook(io.BytesIO(export.content)).active
    assert sheet is not None
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0][2] == "Номер машины"
    assert [row[2] for row in rows[1:3]] == ["01 A 124 BC", "01 B 777 CD"]
    assert rows[-1][9] == 451

    missing = await client.get(f"/api/loading/sessions/{'c' * 32}")
    assert missing.status_code == 404


async def test_cargo_classes_must_be_detected(client: httpx.AsyncClient) -> None:
    detection = {"model": "yolo26n.pt", "confidence": 0.4, "classes": ["truck", "suitcase"]}
    bad = await client.put(
        "/api/settings", json={"detection": detection | {"cargo_classes": ["backpack"]}}
    )
    assert bad.status_code == 422
    good = await client.put(
        "/api/settings", json={"detection": detection | {"cargo_classes": ["suitcase"]}}
    )
    assert good.status_code == 200
    assert good.json()["detection"]["cargo_classes"] == ["suitcase"]
